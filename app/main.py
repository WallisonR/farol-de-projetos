import csv, hmac, io, json, os, time
from contextlib import asynccontextmanager
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from fastapi import FastAPI, Depends, HTTPException, Header, UploadFile, File, Query, APIRouter
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError, ProgrammingError
from sqlalchemy.orm import Session, joinedload
from . import rules as R, auth as A
from .db import Base, engine, migrar, SessionLocal, SEM_BANCO, URL, Conta, Projeto, Historico, Importacao, User
from .seed import seed

PUBLIC = Path(__file__).parent.parent / "public"
_pronto = False
ADMIN_EMAIL = lambda: os.getenv("ADMIN_EMAIL", "wallison@chatbotmaker.io").strip().lower()  # noqa: E731


def iniciar():
    """Cria tabelas, migra colunas, garante o admin e faz a carga inicial (idempotente, seguro em serverless)."""
    global _pronto
    if _pronto:
        return
    try:
        Base.metadata.create_all(engine)
    except (ProgrammingError, IntegrityError):
        pass
    migrar()
    with SessionLocal() as db:
        try:
            adm = db.scalar(select(User).where(User.email == ADMIN_EMAIL()))
            senha = os.getenv("ADMIN_SENHA")
            if not adm and senha:
                adm = User(email=ADMIN_EMAIL(), nome=os.getenv("ADMIN_NOME", "Wallison Rocha"),
                           senha_hash=A.hash_senha(senha), papel="admin")
                db.add(adm); db.flush()
            if adm:
                if os.getenv("SEED", "1") != "0":
                    seed(db, adm.id)
                db.execute(update(Conta).where(Conta.responsavel_id.is_(None)).values(responsavel_id=adm.id))
                db.execute(update(Historico).where(Historico.responsavel_id.is_(None)).values(responsavel_id=adm.id))
            db.commit()
        except IntegrityError:
            db.rollback()
    _pronto = True


@asynccontextmanager
async def lifespan(app):
    if not SEM_BANCO:
        iniciar()
    yield


app = FastAPI(title="Farol da Carteira · Suri Shop by TOTVS", version="2.0", lifespan=lifespan)


def get_db():
    if SEM_BANCO:
        raise HTTPException(503, "Banco não configurado: defina DATABASE_URL (PostgreSQL) nas variáveis de ambiente do Vercel")
    iniciar()
    with SessionLocal() as db:
        yield db


def ator(authorization: str | None = Header(None), db: Session = Depends(get_db)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Faça login")
    tok = authorization[7:]
    svc = os.getenv("API_TOKEN")
    if svc and hmac.compare_digest(tok, svc):  # token de serviço (webhooks/integrações): acesso de administrador
        return SimpleNamespace(id=None, email="api", nome="API", papel="admin", ativo=True)
    uid = A.ler(tok)
    u = db.get(User, uid) if uid else None
    if not u or not u.ativo:
        raise HTTPException(401, "Sessão inválida ou expirada")
    return u


def so_admin(a=Depends(ator)):
    if a.papel != "admin":
        raise HTTPException(403, "Apenas administradores")
    return a


pub = APIRouter(prefix="/api")
api = APIRouter(prefix="/api", dependencies=[Depends(ator)])
FALHAS: dict[str, list[float]] = {}


# ---------- schemas ----------
class ContaIn(BaseModel):
    nome: str
    nivel: str
    mrr: float | str
    segmento: str | None = None
    nota: str | None = None
    ajuste_impacto: int = 0
    razao_social: str | None = None
    cnpj: str | None = None
    complexidade: str | None = None
    erp: str | None = None
    integracoes: str | None = None
    contato_nome: str | None = None
    contato_email: str | None = None
    contato_telefone: str | None = None
    data_inicio: str | None = None
    ultimo_contato: str | None = None
    proximo_contato: str | None = None
    responsavel_id: int | None = None


class ContaPatch(ContaIn):
    nome: str | None = None
    nivel: str | None = None
    mrr: float | str | None = None
    ajuste_impacto: int | None = None
    churn: bool | None = None


class ProjetoIn(BaseModel):
    conta_id: int | None = None
    conta: str | None = None
    nome: str
    tipo: str
    situacao: str | None = None
    critico: bool = False
    bloqueio_dono: str | None = None
    bloqueio_motivo: str | None = None
    prazo: str | None = None
    proxima_acao: str | None = None
    ajuste_impacto: int = 0


class ProjetoPatch(BaseModel):
    nome: str | None = None
    tipo: str | None = None
    situacao: str | None = None
    critico: bool | None = None
    bloqueio_dono: str | None = None
    bloqueio_motivo: str | None = None
    prazo: str | None = None
    proxima_acao: str | None = None
    ajuste_impacto: int | None = None


class Evento(ProjetoPatch):
    conta: str
    projeto: str


class Login(BaseModel):
    email: str
    senha: str


class NovaSenha(BaseModel):
    atual: str
    nova: str


class UserIn(BaseModel):
    email: str
    nome: str
    senha: str
    papel: str = "am"


class UserPatch(BaseModel):
    nome: str | None = None
    papel: str | None = None
    ativo: bool | None = None
    senha: str | None = None


# ---------- helpers ----------
def falha(e, w=()):
    raise HTTPException(422, detail={"erros": list(e), "avisos": list(w)})


def q_contas(a, resp_id=None):
    q = select(Conta)
    if a.papel != "admin":
        q = q.where(Conta.responsavel_id == a.id)
    elif resp_id:
        q = q.where(Conta.responsavel_id == resp_id)
    return q


def lookup_contas(db, a=None, excluir=None):
    """a=None: todas as contas (unicidade global); com ator: só as visíveis (cadastro de projetos)."""
    q = q_contas(a) if a is not None else select(Conta)
    return {c.nome_norm: c.nome for c in db.scalars(q) if c.id != excluir}


def existentes(db):
    rows = db.execute(select(Conta.nome_norm, Projeto.nome_norm).join(Projeto, Projeto.conta_id == Conta.id)).all()
    return {(x, y) for x, y in rows}


def dono_para(db, a, pedido):
    if a.papel != "admin":
        return a.id
    if pedido:
        if not db.get(User, pedido):
            falha(["Responsável não encontrado"])
        return pedido
    if a.id:
        return a.id
    adm = db.scalar(select(User).where(User.papel == "admin", User.ativo == True).order_by(User.id))  # noqa: E712
    return adm.id if adm else None


def log(db, a, ent, eid, rotulo, campo, ant, novo, dono):
    db.add(Historico(usuario=a.email, entidade=ent, entidade_id=eid, rotulo=rotulo, campo=campo, responsavel_id=dono,
                     valor_anterior=None if ant is None else str(ant), valor_novo=None if novo is None else str(novo)))


CAMPOS_CONTA = ("razao_social", "cnpj", "complexidade", "erp", "integracoes", "contato_nome", "contato_email",
                "contato_telefone", "data_inicio", "ultimo_contato", "proximo_contato")


def _d(v):
    return v.isoformat() if v else None


def c_out(c, n=None, users=None):
    o = dict(id=c.id, nome=c.nome, nivel=c.nivel, mrr=c.mrr, segmento=c.segmento, nota=c.nota, churn=c.churn,
             ajuste_impacto=c.ajuste_impacto, n_projetos=n, responsavel_id=c.responsavel_id,
             responsavel=(users or {}).get(c.responsavel_id))
    for k in CAMPOS_CONTA:
        v = getattr(c, k)
        o[k] = v.isoformat() if isinstance(v, date) else v
    return o


def p_out(p, hoje):
    cor, mot = R.farol(p, hoje)
    return dict(id=p.id, conta_id=p.conta_id, conta=p.conta.nome, nome=p.nome, tipo=p.tipo, situacao=p.situacao,
                critico=p.critico, bloqueio_dono=p.bloqueio_dono, bloqueio_motivo=p.bloqueio_motivo,
                prazo=_d(p.prazo), proxima_acao=p.proxima_acao, ajuste_impacto=p.ajuste_impacto,
                ultima_atualizacao=_d(p.ultima_atualizacao), farol=dict(cor=cor, motivo=mot),
                prioridade=R.prioridade(p, p.conta, hoje))


def nomes_users(db):
    return {u.id: u.nome for u in db.scalars(select(User))}


def projetos_ativos(db, a, resp_id=None):
    q = select(Projeto).join(Conta).options(joinedload(Projeto.conta)).where(Conta.churn == False)  # noqa: E712
    if a.papel != "admin":
        q = q.where(Conta.responsavel_id == a.id)
    elif resp_id:
        q = q.where(Conta.responsavel_id == resp_id)
    return list(db.scalars(q).unique())


def conta_visivel(db, a, cid):
    c = db.get(Conta, cid)
    if not c or (a.papel != "admin" and c.responsavel_id != a.id):
        raise HTTPException(404, "Conta não encontrada")
    return c


def projeto_visivel(db, a, pid):
    p = db.get(Projeto, pid)
    if not p or (a.papel != "admin" and p.conta.responsavel_id != a.id):
        raise HTTPException(404, "Projeto não encontrado")
    return p


def criar_conta(db, l, a, dono):
    c = Conta(**l, responsavel_id=dono)
    db.add(c); db.flush()
    log(db, a, "conta", c.id, c.nome, "conta", None, "cadastrada", dono)
    return c


def criar_projeto(db, l, a):
    conta = db.scalar(select(Conta).where(Conta.nome_norm == l["conta_norm"]))
    p = Projeto(conta_id=conta.id, nome=l["nome"], nome_norm=l["nome_norm"], tipo=l["tipo"], situacao=l["situacao"],
                critico=l["critico"], bloqueio_dono=l["bloqueio_dono"], bloqueio_motivo=l["bloqueio_motivo"],
                prazo=l["prazo"], proxima_acao=l["proxima_acao"], ajuste_impacto=l["ajuste_impacto"],
                ultima_atualizacao=date.today())
    db.add(p); db.flush()
    log(db, a, "projeto", p.id, f"{conta.nome} / {p.nome}", "projeto", None, "cadastrado", conta.responsavel_id)
    return p


def atualizar_projeto(db, p, mudancas, a):
    cur = dict(conta=p.conta.nome, nome=p.nome, tipo=p.tipo, situacao=p.situacao, critico=p.critico,
               bloqueio_dono=p.bloqueio_dono, bloqueio_motivo=p.bloqueio_motivo, prazo=p.prazo,
               proxima_acao=p.proxima_acao, ajuste_impacto=p.ajuste_impacto)
    cur.update(mudancas)
    ex = existentes(db) - {(R.norm(p.conta.nome), p.nome_norm)}
    e, w, l = R.valida_projeto(cur, {R.norm(p.conta.nome): p.conta.nome}, ex)
    if e:
        falha(e, w)
    hoje = date.today()
    antes = R.farol(p, hoje)[0]
    rotulo, dono = f"{p.conta.nome} / {p.nome}", p.conta.responsavel_id
    for k in ("nome", "nome_norm", "tipo", "situacao", "critico", "bloqueio_dono", "bloqueio_motivo", "prazo",
              "proxima_acao", "ajuste_impacto"):
        if getattr(p, k) != l[k]:
            if k != "nome_norm":
                log(db, a, "projeto", p.id, rotulo, k, getattr(p, k), l[k], dono)
            setattr(p, k, l[k])
    p.ultima_atualizacao = hoje
    depois = R.farol(p, hoje)[0]
    if antes != depois:
        log(db, a, "projeto", p.id, rotulo, "farol", antes, depois, dono)
    return w


def u_out(u):
    return dict(id=u.id, email=u.email, nome=u.nome, papel=u.papel, ativo=u.ativo)


# ---------- autenticação e usuários ----------
@pub.post("/auth/login")
def login(body: Login, db: Session = Depends(get_db)):
    email = body.email.strip().lower()
    agora = time.time()
    tent = [t for t in FALHAS.get(email, []) if agora - t < 900]
    if len(tent) >= 5:
        raise HTTPException(429, "Muitas tentativas. Aguarde 15 minutos.")
    u = db.scalar(select(User).where(User.email == email))
    if not u or not u.ativo or not A.confere(body.senha, u.senha_hash):
        FALHAS[email] = tent + [agora]
        raise HTTPException(401, "E-mail ou senha inválidos")
    FALHAS.pop(email, None)
    return dict(token=A.emitir(u.id), usuario=u_out(u))


@api.get("/auth/eu")
def eu(a=Depends(ator)):
    return dict(id=a.id, email=a.email, nome=a.nome, papel=a.papel)


@api.post("/auth/senha")
def trocar_senha(body: NovaSenha, a=Depends(ator), db: Session = Depends(get_db)):
    u = db.get(User, a.id) if a.id else None
    if not u or not A.confere(body.atual, u.senha_hash):
        falha(["Senha atual incorreta"])
    if len(body.nova) < 8:
        falha(["A nova senha precisa ter ao menos 8 caracteres"])
    u.senha_hash = A.hash_senha(body.nova)
    db.commit()
    return {"ok": True}


@api.get("/usuarios")
def listar_usuarios(db: Session = Depends(get_db), a=Depends(so_admin)):
    return [u_out(u) for u in db.scalars(select(User).order_by(User.nome))]


@api.post("/usuarios", status_code=201)
def criar_usuario(body: UserIn, db: Session = Depends(get_db), a=Depends(so_admin)):
    email, e = body.email.strip().lower(), []
    if "@" not in email or "." not in email.split("@")[-1]:
        e.append("E-mail inválido")
    if not body.nome.strip():
        e.append("Nome obrigatório")
    if len(body.senha) < 8:
        e.append("A senha provisória precisa ter ao menos 8 caracteres")
    if body.papel not in ("admin", "am"):
        e.append("Papel deve ser admin ou am")
    if db.scalar(select(User).where(User.email == email)):
        e.append("E-mail já cadastrado")
    if e:
        falha(e)
    u = User(email=email, nome=body.nome.strip(), senha_hash=A.hash_senha(body.senha), papel=body.papel)
    db.add(u); db.commit()
    return u_out(u)


@api.patch("/usuarios/{uid}")
def editar_usuario(uid: int, body: UserPatch, db: Session = Depends(get_db), a=Depends(so_admin)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "Usuário não encontrado")
    d = body.model_dump(exclude_unset=True)
    if uid == a.id and (d.get("ativo") is False or d.get("papel") == "am"):
        falha(["Você não pode desativar nem rebaixar a si mesmo"])
    if "papel" in d and d["papel"] not in ("admin", "am"):
        falha(["Papel deve ser admin ou am"])
    if d.get("senha"):
        if len(d["senha"]) < 8:
            falha(["A senha precisa ter ao menos 8 caracteres"])
        u.senha_hash = A.hash_senha(d["senha"])
    for k in ("nome", "papel", "ativo"):
        if d.get(k) is not None:
            setattr(u, k, d[k])
    db.commit()
    return u_out(u)


# ---------- contas ----------
@api.get("/contas")
def listar_contas(db: Session = Depends(get_db), a=Depends(ator), q: str | None = None, nivel: str | None = None,
                  segmento: str | None = None, complexidade: str | None = None, responsavel_id: int | None = None,
                  churn: bool | None = None):
    cont = {}
    for p in db.scalars(select(Projeto)):
        cont[p.conta_id] = cont.get(p.conta_id, 0) + 1
    users = nomes_users(db)
    cs = db.scalars(q_contas(a, responsavel_id).order_by(Conta.mrr.desc()))
    out = [c_out(c, cont.get(c.id, 0), users) for c in cs]
    if nivel:
        out = [c for c in out if c["nivel"] == nivel]
    if segmento:
        out = [c for c in out if R.norm(segmento) == R.norm(c["segmento"] or "")]
    if complexidade:
        out = [c for c in out if (c["complexidade"] or "") == complexidade]
    if churn is not None:
        out = [c for c in out if c["churn"] == churn]
    if q:
        out = [c for c in out if R.norm(q) in R.norm(" ".join(str(c[k] or "") for k in ("nome", "segmento", "contato_nome", "erp")))]
    return out


@api.post("/contas", status_code=201)
def cadastrar_conta(body: ContaIn, db: Session = Depends(get_db), a=Depends(ator)):
    d = body.model_dump()
    e, w, l = R.valida_conta(d, lookup_contas(db))
    if e:
        falha(e, w)
    c = criar_conta(db, l, a, dono_para(db, a, d.get("responsavel_id")))
    db.commit()
    return dict(c_out(c, 0, nomes_users(db)), avisos=w)


@api.get("/contas/{cid}")
def detalhe_conta(cid: int, db: Session = Depends(get_db), a=Depends(ator)):
    c = conta_visivel(db, a, cid)
    hoje = date.today()
    ps = [p_out(p, hoje) for p in db.scalars(select(Projeto).where(Projeto.conta_id == cid).options(joinedload(Projeto.conta)))]
    return dict(c_out(c, len(ps), nomes_users(db)), projetos=ps)


@api.patch("/contas/{cid}")
def editar_conta(cid: int, body: ContaPatch, db: Session = Depends(get_db), a=Depends(ator)):
    c = conta_visivel(db, a, cid)
    cur = dict(nome=c.nome, nivel=c.nivel, mrr=c.mrr, segmento=c.segmento, nota=c.nota,
               ajuste_impacto=c.ajuste_impacto, churn=c.churn, **{k: getattr(c, k) for k in CAMPOS_CONTA})
    mud = body.model_dump(exclude_unset=True)
    novo_dono = mud.pop("responsavel_id", None) if a.papel == "admin" else None
    cur.update(mud)
    e, w, l = R.valida_conta(cur, lookup_contas(db, excluir=cid))
    if e:
        falha(e, w)
    for k, v in l.items():
        if getattr(c, k) != v:
            if k != "nome_norm":
                log(db, a, "conta", c.id, c.nome, k, getattr(c, k), v, c.responsavel_id)
            setattr(c, k, v)
    if novo_dono and novo_dono != c.responsavel_id:
        if not db.get(User, novo_dono):
            falha(["Responsável não encontrado"])
        log(db, a, "conta", c.id, c.nome, "responsavel", c.responsavel_id, novo_dono, novo_dono)
        c.responsavel_id = novo_dono
    db.commit()
    return dict(c_out(c, None, nomes_users(db)), avisos=w)


# ---------- projetos ----------
@api.get("/projetos")
def listar_projetos(db: Session = Depends(get_db), a=Depends(ator), conta_id: int | None = None, tipo: str | None = None,
                    farol: str | None = None, q: str | None = None, bloqueado: bool | None = None):
    hoje = date.today()
    out = [p_out(p, hoje) for p in projetos_ativos(db, a)]
    if conta_id:
        out = [p for p in out if p["conta_id"] == conta_id]
    if tipo:
        out = [p for p in out if p["tipo"] == tipo]
    if farol:
        out = [p for p in out if p["farol"]["cor"] == farol.upper()]
    if bloqueado is not None:
        out = [p for p in out if bool(p["bloqueio_dono"]) == bloqueado]
    if q:
        out = [p for p in out if R.norm(q) in R.norm(p["conta"] + " " + p["nome"])]
    return sorted(out, key=lambda p: -p["prioridade"]["ordem"])


@api.post("/projetos", status_code=201)
def cadastrar_projeto(body: ProjetoIn, db: Session = Depends(get_db), a=Depends(ator)):
    d = body.model_dump()
    if d["conta_id"]:
        c = db.get(Conta, d["conta_id"])
        d["conta"] = c.nome if c and (a.papel == "admin" or c.responsavel_id == a.id) else None
    e, w, l = R.valida_projeto(d, lookup_contas(db, a), existentes(db))
    if e:
        falha(e, w)
    p = criar_projeto(db, l, a)
    db.commit()
    db.refresh(p)
    return dict(p_out(p, date.today()), avisos=w)


@api.get("/projetos/{pid}")
def detalhe_projeto(pid: int, db: Session = Depends(get_db), a=Depends(ator)):
    p = projeto_visivel(db, a, pid)
    h = db.scalars(select(Historico).where(Historico.entidade == "projeto", Historico.entidade_id == pid)
                   .order_by(Historico.id.desc()).limit(100))
    return dict(p_out(p, date.today()), historico=[hist_out(x) for x in h])


@api.patch("/projetos/{pid}")
def editar_projeto(pid: int, body: ProjetoPatch, db: Session = Depends(get_db), a=Depends(ator)):
    p = projeto_visivel(db, a, pid)
    w = atualizar_projeto(db, p, body.model_dump(exclude_unset=True), a)
    db.commit()
    return dict(p_out(p, date.today()), avisos=w)


# ---------- carteira e histórico ----------
@api.get("/carteira/farol")
def carteira(db: Session = Depends(get_db), a=Depends(ator), responsavel_id: int | None = None):
    hoje = date.today()
    users = nomes_users(db)
    ps = sorted((p_out(p, hoje) for p in projetos_ativos(db, a, responsavel_id)), key=lambda p: -p["prioridade"]["ordem"])
    contas = [c_out(c, sum(1 for p in ps if p["conta_id"] == c.id), users)
              for c in db.scalars(q_contas(a, responsavel_id).order_by(Conta.mrr.desc()))]
    ativas = [c for c in contas if not c["churn"]]
    cnt = lambda cor: sum(1 for p in ps if p["farol"]["cor"] == cor)  # noqa: E731
    dono = {}
    for p in ps:
        if p["bloqueio_dono"]:
            dono[p["bloqueio_dono"]] = dono.get(p["bloqueio_dono"], 0) + 1
    resumo = dict(contas_ativas=len(ativas), projetos=len(ps), vermelho=cnt("R"), amarelo=cnt("Y"), verde=cnt("G"),
                  cinza=cnt("C"), bloqueados=sum(1 for p in ps if p["bloqueio_dono"]),
                  atrasados=sum(1 for p in ps if p["prazo"] and p["prazo"] < hoje.isoformat()),
                  mrr_ativo=round(sum(c["mrr"] for c in ativas), 2), bloqueios_por_dono=dono)
    resp = [dict(id=i, nome=n) for i, n in users.items()] if a.papel == "admin" else []
    return dict(hoje=hoje.isoformat(), resumo=resumo, contas=contas, projetos=ps, responsaveis=resp)


def hist_out(h):
    return dict(id=h.id, em=h.em.isoformat(), usuario=h.usuario, entidade=h.entidade, rotulo=h.rotulo, campo=h.campo,
                valor_anterior=h.valor_anterior, valor_novo=h.valor_novo)


@api.get("/historico")
def historico(db: Session = Depends(get_db), a=Depends(ator), limit: int = Query(200, le=1000)):
    q = select(Historico).order_by(Historico.id.desc()).limit(limit)
    if a.papel != "admin":
        q = q.where(Historico.responsavel_id == a.id)
    return [hist_out(h) for h in db.scalars(q)]


# ---------- importação ----------
MC = {"nome": ["nome", "conta", "cliente"], "nivel": ["nivel", "tier"], "mrr": ["mrr", "receita"],
      "segmento": ["segmento"], "nota": ["observacoes", "obs", "nota"], "razao_social": ["razaosocial"],
      "cnpj": ["cnpj"], "complexidade": ["complexidade"], "erp": ["erp", "sistema"], "integracoes": ["integracoes"],
      "contato_nome": ["contato", "contatonome", "nomedocontato"], "contato_email": ["email", "emaildocontato", "contatoemail"],
      "contato_telefone": ["telefone", "whatsapp", "contatotelefone"], "data_inicio": ["inicio", "iniciodocontrato", "datadeinicio"],
      "ultimo_contato": ["ultimocontato"], "proximo_contato": ["proximocontato"]}
MP = {"conta": ["conta", "cliente"], "nome": ["projeto", "nomedoprojeto"], "tipo": ["tipo"],
      "situacao": ["situacao", "status"], "prazo": ["prazo"], "proxima_acao": ["proximaacao"],
      "bloqueio_dono": ["bloqueio"], "bloqueio_motivo": ["motivodobloqueio", "motivo"], "critico": ["critico"]}


def mapear(row, M):
    o = {}
    for k, v in row.items():
        h = R.hk(k)
        for f, nomes in M.items():
            if h in nomes and f not in o:
                o[f] = v
    if "critico" in o:
        o["critico"] = R.hk(o["critico"]) in ("sim", "s", "true", "x", "verdadeiro") or str(o["critico"]).strip() == "1"
    return o


def _tipo(rows):
    return "p" if any(R.hk(k) == "projeto" for k in rows[0]) else "c"


def ler_arquivo(nome, dados):
    n = nome.lower()
    try:
        if n.endswith(".json"):
            j = json.loads(dados.decode("utf-8-sig"))
            return [("c", j.get("contas", [])), ("p", j.get("projetos", []))]
        if n.endswith(".csv"):
            try:
                t = dados.decode("utf-8-sig")
            except UnicodeDecodeError:
                t = dados.decode("latin-1")
            try:
                dial = csv.Sniffer().sniff(t[:2000], delimiters=",;\t")
            except csv.Error:
                dial = csv.excel
            rows = list(csv.DictReader(io.StringIO(t), dialect=dial))
            return [(_tipo(rows), rows)] if rows else []
        if n.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(dados), data_only=True, read_only=True)
            out = []
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                hdr = next(it, None)
                if not hdr:
                    continue
                rows = [{str(h): v for h, v in zip(hdr, r) if h is not None} for r in it
                        if any(x not in (None, "") for x in r)]
                if rows:
                    out.append((_tipo(rows), rows))
            return out
    except HTTPException:
        raise
    except Exception as ex:
        raise HTTPException(400, f"Não foi possível ler o arquivo: {ex}")
    raise HTTPException(400, "Formato não suportado. Use .xlsx, .csv ou .json (salve .xls como .xlsx)")


def processar(db, a, sets):
    lk_todas, lk_proj, ex, itens = lookup_contas(db), lookup_contas(db, a), existentes(db), []
    for kind, rows in sorted(sets, key=lambda s: s[0]):  # contas antes de projetos
        for i, r in enumerate(rows, 2):
            if kind == "c":
                e, w, l = R.valida_conta(mapear(r, MC), lk_todas)
                if not e:
                    lk_todas[l["nome_norm"]] = lk_proj[l["nome_norm"]] = l["nome"]
                itens.append(dict(tipo="conta", linha=i, registro=l["nome"] or "?", erros=e, avisos=w, _l=l))
            else:
                o = mapear(r, MP)
                e, w, l = R.valida_projeto(o, lk_proj, ex)
                if not e:
                    ex.add((l["conta_norm"], l["nome_norm"]))
                itens.append(dict(tipo="projeto", linha=i, registro=f'{l["conta_nome"] or o.get("conta") or "?"} / {l["nome"] or "?"}',
                                  erros=e, avisos=w, _l=l))
    return itens


@api.post("/importacoes")
async def importar(arquivo: UploadFile = File(...), dry_run: bool = False, db: Session = Depends(get_db), a=Depends(ator)):
    sets = ler_arquivo(arquivo.filename or "", await arquivo.read())
    itens = processar(db, a, sets)
    ok = [x for x in itens if not x["erros"]]
    if not itens:
        raise HTTPException(400, "Nenhuma linha encontrada no arquivo")
    if not dry_run and ok:
        dono = dono_para(db, a, None)
        for x in ok:
            if x["tipo"] == "conta":
                criar_conta(db, x["_l"], a, dono)
        db.flush()
        for x in ok:
            if x["tipo"] == "projeto":
                criar_projeto(db, x["_l"], a)
        db.add(Importacao(arquivo=arquivo.filename or "", linhas_ok=len(ok), linhas_erro=len(itens) - len(ok), usuario=a.email))
        db.commit()
    for x in itens:
        x.pop("_l")
    return dict(arquivo=arquivo.filename, dry_run=dry_run, gravados=0 if dry_run else len(ok), validos=len(ok),
                com_erro=len(itens) - len(ok), linhas=itens)


# ---------- webhook ----------
@api.post("/webhooks/projetos")
def webhook(ev: Evento, db: Session = Depends(get_db), a=Depends(ator)):
    ck, pk = R.norm(ev.conta), R.norm(ev.projeto)
    conta = db.scalar(q_contas(a).where(Conta.nome_norm == ck))
    if not conta:
        falha(["Conta não cadastrada: cadastre a conta antes"])
    p = db.scalar(select(Projeto).where(Projeto.conta_id == conta.id, Projeto.nome_norm == pk))
    mud = ev.model_dump(exclude_unset=True, exclude={"conta", "projeto"})
    if p:
        w = atualizar_projeto(db, p, mud, a)
        acao = "atualizado"
    else:
        e, w, l = R.valida_projeto(dict(mud, conta=conta.nome, nome=ev.projeto), {ck: conta.nome}, existentes(db))
        if e:
            falha(e, w)
        p = criar_projeto(db, l, a)
        acao = "criado"
    db.commit()
    return dict(acao=acao, projeto=p_out(p, date.today()), avisos=w)


@pub.get("/health", include_in_schema=False)
def health():
    return {"ok": not SEM_BANCO, "banco": URL.split(":")[0], "configurado": not SEM_BANCO,
            "admin_senha_definida": bool(os.getenv("ADMIN_SENHA"))}


app.include_router(pub)
app.include_router(api)


@app.get("/", include_in_schema=False)
def index():
    return FileResponse(PUBLIC / "index.html")
